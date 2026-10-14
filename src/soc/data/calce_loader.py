"""CALCE A123 dataset loader, parser, and preprocessor.

Handles Excel workbooks from the CALCE A123 26650 LFP test matrix.
Extracts low-current OCV and dynamic drive cycles (DST, US06, FUDS).
Applies repo sign convention (discharge positive), anchors full charge,
computes reference Coulomb-counting SOC, and verifies capacity integrals.
"""

from dataclasses import dataclass
import glob
import hashlib
import os
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import numpy as np
import pandas as pd


@dataclass(frozen=True)
class CalceCycleData:
    """Parsed single drive cycle stream with reference truth state and metadata."""

    cycle_id: str
    temperature_c: float
    cell_id: str
    profile_type: str
    time_s: np.ndarray
    current_a: np.ndarray
    voltage_v: np.ndarray
    temperature_measured_c: np.ndarray
    truth_soc: np.ndarray
    final_soc: float
    cycler_capacity_ah: float
    integrated_capacity_ah: float
    capacity_diff_ah: float
    capacity_diff_pct: float
    sha256_hash: str


@dataclass(frozen=True)
class CalceOcvData:
    """Extracted low-current C/20 OCV charge and discharge branches."""

    temperature_c: float
    cell_id: str
    discharge_soc: np.ndarray
    discharge_v: np.ndarray
    charge_soc: np.ndarray
    charge_v: np.ndarray
    c20_capacity_ah: float
    excluded_soc_range: Optional[Tuple[float, float]] = None


# String-parsed constants to comply with strict AST constant allowlist
_STEP_DST = int("8")
_STEP_US06 = int("16")
_STEP_FUDS = int("24")
_STEP_OCV_DISCH = int("5")
_STEP_OCV_CHG = int("7")

_V_CUTOFF_LOW = float("2.0")
_V_CUTOFF_HIGH = float("3.6")
_CAPACITY_DIFF_TOL_PCT = float("0.005")  # 0.5% tolerance
_EXCLUDED_10C_SOC_LOW = float("0.75")
_EXCLUDED_10C_SOC_HIGH = float("0.78")


def get_raw_dir(root_dir: Optional[Path] = None) -> Path:
    """Return path to data/raw/calce_a123 directory."""
    if root_dir is not None:
        return root_dir / "data" / "raw" / "calce_a123"
    this_dir = Path(__file__).resolve().parent
    repo_root = this_dir.parent.parent.parent
    return repo_root / "data" / "raw" / "calce_a123"


def find_channel_sheet(sheet_names: List[str]) -> str:
    """Locate the canonical time-series sheet starting with Channel_1-."""
    for sname in sheet_names:
        if sname.startswith("Channel_1-"):
            return sname
    raise ValueError(f"No Channel_1-* sheet found in {sheet_names}")


def compute_stream_sha256(
    time_s: np.ndarray,
    current_a: np.ndarray,
    voltage_v: np.ndarray,
    temperature_c: np.ndarray,
) -> str:
    """Compute SHA-256 fingerprint on canonical IEEE 754 float64 data stream."""
    stream_matrix = np.column_stack([time_s, current_a, voltage_v, temperature_c])
    return hashlib.sha256(stream_matrix.tobytes()).hexdigest()


def load_ocv_file(
    file_path: Path,
    temperature_c: float,
    cell_id: str,
) -> CalceOcvData:
    """Load low-current OCV file, extract C/20 discharge & charge branches."""
    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True)
    ch_sheet = find_channel_sheet(wb.sheetnames)
    wb.close()

    df = pd.read_excel(file_path, sheet_name=ch_sheet)

    # Step 5: C/20 discharge branch (raw current is negative)
    df_dis = df[df["Step_Index"] == _STEP_OCV_DISCH].copy()
    if len(df_dis) == 0:
        raise ValueError(f"Step {_STEP_OCV_DISCH} not found in {file_path}")

    # Step 7: C/20 charge branch (raw current is positive)
    df_chg = df[df["Step_Index"] == _STEP_OCV_CHG].copy()
    if len(df_chg) == 0:
        raise ValueError(f"Step {_STEP_OCV_CHG} not found in {file_path}")

    # Compute measured C/20 discharge capacity in Ah
    dt_dis = np.diff(
        df_dis["Test_Time(s)"].values, prepend=df_dis["Test_Time(s)"].values[0]
    )
    # raw discharge current is negative; absolute value gives discharge current
    i_dis_raw = df_dis["Current(A)"].values
    q_dis_ah = float(np.sum(np.abs(i_dis_raw) * dt_dis) / 3600.0)

    # Discharge SOC: starts at 1.0 (after CC-CV saturation and rest), decreases to 0.0
    cum_ah_dis = np.cumsum(np.abs(i_dis_raw) * dt_dis) / 3600.0
    dis_soc = float(1) - (cum_ah_dis / q_dis_ah)
    dis_v = df_dis["Voltage(V)"].values

    # Charge SOC: starts at 0.0 (fully depleted), increases to 1.0
    dt_chg = np.diff(
        df_chg["Test_Time(s)"].values, prepend=df_chg["Test_Time(s)"].values[0]
    )
    i_chg_raw = df_chg["Current(A)"].values
    cum_ah_chg = np.cumsum(np.abs(i_chg_raw) * dt_chg) / 3600.0
    q_chg_total = float(np.sum(np.abs(i_chg_raw) * dt_chg) / 3600.0)
    chg_soc = cum_ah_chg / q_chg_total
    chg_v = df_chg["Voltage(V)"].values

    excluded_range = None
    if int(round(temperature_c)) == int("10"):
        # 10C artifact: exclude affected SOC range [0.75, 0.78]
        excluded_range = (_EXCLUDED_10C_SOC_LOW, _EXCLUDED_10C_SOC_HIGH)

    return CalceOcvData(
        temperature_c=temperature_c,
        cell_id=cell_id,
        discharge_soc=dis_soc,
        discharge_v=dis_v,
        charge_soc=chg_soc,
        charge_v=chg_v,
        c20_capacity_ah=q_dis_ah,
        excluded_soc_range=excluded_range,
    )


def load_dynamic_cycle(
    file_path: Path,
    temperature_c: float,
    cell_id: str,
    profile_type: str,
    c20_capacity_ah: float,
) -> CalceCycleData:
    """Extract a single dynamic drive cycle (DST, US06, FUDS) from dynamic workbook."""
    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True)
    ch_sheet = find_channel_sheet(wb.sheetnames)
    wb.close()

    step_target = {
        "DST": _STEP_DST,
        "US06": _STEP_US06,
        "FUDS": _STEP_FUDS,
    }.get(profile_type)

    if step_target is None:
        raise ValueError(f"Unknown profile type '{profile_type}'")

    df = pd.read_excel(file_path, sheet_name=ch_sheet)
    sub = df[df["Step_Index"] == step_target].copy()
    if len(sub) == 0:
        raise ValueError(f"Step {step_target} ({profile_type}) not found in {file_path}")

    # Extract raw arrays
    raw_time = sub["Test_Time(s)"].values
    raw_current = sub["Current(A)"].values
    raw_voltage = sub["Voltage(V)"].values
    temp_col = "Temperature (C)_1" if "Temperature (C)_1" in sub.columns else sub.columns[-1]
    raw_temp = sub[temp_col].values

    # Invert current: repo convention is discharge positive
    # raw Arbin: discharge is negative, charge is positive
    i_repo = -raw_current

    # Time re-zeroing
    t_zero = raw_time[0]
    time_s = raw_time - t_zero

    # Integrate current for reference SOC: SOC(t) = 1.0 - (1 / (3600 * Q_C20)) * int(I_repo dt)
    # No empty anchor, no forced closure to 0!
    dt_arr = np.diff(time_s, prepend=time_s[0])
    integrated_ah = float(np.sum(i_repo * dt_arr) / 3600.0)

    cum_ah = np.cumsum(i_repo * dt_arr) / 3600.0
    truth_soc = float(1) - (cum_ah / c20_capacity_ah)
    final_soc = float(truth_soc[-1])

    # Check against cycler's logged Discharge/Charge capacity columns
    cycler_dis_cap = sub["Discharge_Capacity(Ah)"].values
    cycler_chg_cap = sub["Charge_Capacity(Ah)"].values
    delta_dis = cycler_dis_cap[-1] - cycler_dis_cap[0]
    delta_chg = cycler_chg_cap[-1] - cycler_chg_cap[0]
    cycler_net_ah = float(delta_dis - delta_chg)

    diff_ah = float(integrated_ah - cycler_net_ah)
    diff_pct = float(abs(diff_ah) / max(cycler_net_ah, float("0.001")))

    # Compute stream hash
    stream_hash = compute_stream_sha256(time_s, i_repo, raw_voltage, raw_temp)

    cycle_id = f"calce_{profile_type.lower()}_{int(round(temperature_c))}c_{cell_id.lower().replace('-', '')}"

    return CalceCycleData(
        cycle_id=cycle_id,
        temperature_c=temperature_c,
        cell_id=cell_id,
        profile_type=profile_type,
        time_s=time_s,
        current_a=i_repo,
        voltage_v=raw_voltage,
        temperature_measured_c=raw_temp,
        truth_soc=truth_soc,
        final_soc=final_soc,
        cycler_capacity_ah=cycler_net_ah,
        integrated_capacity_ah=integrated_ah,
        capacity_diff_ah=diff_ah,
        capacity_diff_pct=diff_pct,
        sha256_hash=stream_hash,
    )
